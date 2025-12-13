from fastapi import FastAPI, Depends, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select
from contextlib import asynccontextmanager
from typing import List, Optional
from io import BytesIO
import re
import textwrap
import unicodedata
import json
import os
import tempfile
import zipfile

from .database import create_db_and_tables, get_session
from .models import (
    Project,
    Phase,
    Artifact,
    PhaseType,
    ArtifactType,
    CodeGenSession,
    CodeGenFile,
    CodeGenMessage,
)
from .schemas import (
    ProjectCreate, ProjectResponse, ProjectDetailResponse,
    PhaseResponse, ArtifactResponse,
    GenerateRequest, GenerateResponse,
    ExportCustomRequest,
    CodeZipRequest,
    CodeGenCreateSessionRequest,
    CodeGenSessionResponse,
    CodeGenSessionDetailResponse,
    CodeGenUpsertFileRequest,
    CodeGenGenerateFileRequest,
    CodeGenFileResponse,
    CodeGenPlanResponse,
    CodeGenApplyPlanResponse,
    ChatRequest,
    ChatResponse,
)
from .services import (
    create_project_phases, generate_phase_artifacts,
    get_phase_by_type, get_all_artifacts,
    ensure_project_phases,
)
from .llm import check_llm_status, call_coder_model, call_primary_model
from .prompts import SYSTEM_PROMPT
from .codegen import (
    build_generate_file_prompt,
    parse_file_json,
    safe_relpath,
    normalize_content,
    suggest_file_paths,
)
from .config import get_settings
from .guardrails import (
    GuardrailError,
    validate_generated_app_payload,
    validate_single_file_payload,
)


# --- Markdown helpers for export formatting ---

def build_project_markdown(session: Session, project_id: int) -> tuple[Project, str]:
    """Build a single markdown document for a project and all its artifacts."""
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    artifacts = get_all_artifacts(session, project_id)

    markdown = f"# {project.name}\n\n"
    markdown += f"**Description:** {project.description}\n\n"
    markdown += f"**User Type:** {project.user_type.value}\n"
    markdown += f"**Constraints:** {project.constraints or 'None'}\n"
    markdown += f"**Skills:** {project.skills or 'Not specified'}\n"
    markdown += f"**Time Available:** {project.time_available or 'Not specified'}\n\n"
    markdown += "---\n\n"

    current_phase = None
    for artifact in artifacts:
        phase = session.get(Phase, artifact.phase_id)
        if phase and phase.phase_type != current_phase:
            current_phase = phase.phase_type
            markdown += f"## {current_phase.value.replace('_', ' ').title()}\n\n"

        markdown += f"### {artifact.artifact_type.value.replace('_', ' ').title()} (v{artifact.version})\n\n"
        markdown += artifact.content + "\n\n"
        markdown += "---\n\n"

    return project, markdown


def _remove_emojis(text: str) -> str:
    """Best-effort removal of emoji characters so PDFs use only supported glyphs.

    This keeps export robust when using ReportLab's built-in fonts.
    """
    cleaned_chars = []
    for ch in text:
        code = ord(ch)
        # Basic emoji ranges (emoticons, pictographs, transport, flags, etc.)
        if (
            0x1F300 <= code <= 0x1FAFF  # Misc & Supplemental Symbols and Pictographs, Emoji
            or 0x2600 <= code <= 0x26FF  # Misc symbols
            or 0x2700 <= code <= 0x27BF  # Dingbats
        ):
            continue
        # Skip characters explicitly marked as emoji in Unicode name when available
        name = unicodedata.name(ch, "")
        if "EMOJI" in name:
            continue
        cleaned_chars.append(ch)
    return "".join(cleaned_chars)


def clean_inline_markdown(text: str) -> str:
    """Remove common inline markdown markers like **bold**, *italic*, `code`, links, and emojis.

    This is intentionally simple but good enough for export formatting.
    """
    # Strip backticks used for inline code
    text = text.replace("`", "")
    # Bold/italic markers
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"__(.*?)__", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"_(.*?)_", r"\1", text)
    # Links: [label](url) -> label
    text = re.sub(r"\[(.*?)\]\([^)]+\)", r"\1", text)
    # Finally, strip emojis/unsupported pictographs
    text = _remove_emojis(text)
    return text


def iter_markdown_blocks(markdown: str):
    """Very small markdown-to-blocks helper for headings, bullets, and rules.

    Yields (text, kind) where kind is one of: h1, h2, h3, bullet, hr, paragraph, blank.
    """
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip("\n")
        stripped = line.strip()
        if not stripped:
            yield "", "blank"
            continue

        # Horizontal rule markers (rendered as real lines in PDF/DOCX)
        if stripped in {"---", "***", "___"}:
            yield "", "hr"
            continue

        if line.startswith("# "):
            yield clean_inline_markdown(line[2:].strip()), "h1"
        elif line.startswith("## "):
            yield clean_inline_markdown(line[3:].strip()), "h2"
        elif line.startswith("### "):
            yield clean_inline_markdown(line[4:].strip()), "h3"
        elif re.match(r"^\s*[-*+]\s+", line):
            text = re.sub(r"^\s*[-*+]\s+", "", line).strip()
            yield clean_inline_markdown(text), "bullet"
        elif re.match(r"^\s*\d+\.\s+", line):
            m = re.match(r"^\s*(\d+)\.\s+(.*)", line)
            if m:
                numeric = m.group(1)
                body = m.group(2).strip()
                yield clean_inline_markdown(f"{numeric}. {body}"), "bullet"
            else:
                yield clean_inline_markdown(stripped), "paragraph"
        else:
            yield clean_inline_markdown(stripped), "paragraph"


def markdown_to_pdf(markdown: str, base_filename: str) -> StreamingResponse:
    """Render markdown-ish text to a simple but structured PDF (headings/bullets)."""
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="PDF export not available on server (missing reportlab dependency)",
        )

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter
    x_margin = 72
    y = height - 72
    line_height = 16

    for text, kind in iter_markdown_blocks(markdown):
        if kind == "blank":
            # Add a little vertical space
            y -= line_height // 2
            continue

        # Render horizontal rules as actual lines
        if kind == "hr":
            if y < 72 + line_height:
                pdf.showPage()
                y = height - 72
            pdf.setLineWidth(0.5)
            pdf.line(x_margin, y, width - x_margin, y)
            y -= line_height
            continue

        # Choose font based on block kind
        if kind == "h1":
            pdf.setFont("Helvetica-Bold", 18)
            wrap_width = 40  # shorter to avoid clipping long titles
        elif kind == "h2":
            pdf.setFont("Helvetica-Bold", 16)
            wrap_width = 60
        elif kind == "h3":
            pdf.setFont("Helvetica-Bold", 14)
            wrap_width = 80
        else:
            pdf.setFont("Helvetica", 11)
            wrap_width = 95

        prefix = ""
        if kind == "bullet":
            prefix = "• "

        # Wrap long lines so they don't get cut off horizontally
        wrapped_lines = textwrap.wrap(text, width=wrap_width) or [""]

        for i, line in enumerate(wrapped_lines):
            if y < 72:
                pdf.showPage()
                y = height - 72
                # Re-apply font after new page
                if kind == "h1":
                    pdf.setFont("Helvetica-Bold", 18)
                elif kind == "h2":
                    pdf.setFont("Helvetica-Bold", 16)
                elif kind == "h3":
                    pdf.setFont("Helvetica-Bold", 14)
                else:
                    pdf.setFont("Helvetica", 11)

            draw_text = f"{prefix}{line}" if i == 0 else f"  {line}" if kind == "bullet" else line
            pdf.drawString(x_margin, y, draw_text[:1000])
            y -= line_height

    pdf.save()
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{base_filename}.pdf"',
        },
    )


def markdown_to_docx(markdown: str, base_filename: str) -> StreamingResponse:
    """Render markdown-ish text to DOCX with headings and bullet lists."""
    try:
        from docx import Document
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="DOCX export not available on server (missing python-docx dependency)",
        )

    document = Document()

    for text, kind in iter_markdown_blocks(markdown):
        if kind == "blank":
            document.add_paragraph("")
        elif kind == "hr":
            # DOCX doesn't have a native horizontal rule helper; use a styled paragraph
            run = document.add_paragraph().add_run("―" * 30)
            run.bold = True
        elif kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "h2":
            document.add_heading(text, level=2)
        elif kind == "h3":
            document.add_heading(text, level=3)
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{base_filename}.docx"',
        },
    )


@asynccontextmanager
async def lifespan(app: FastAPI):
    create_db_and_tables()
    yield


app = FastAPI(
    title="AI Innovation Hub",
    description="Transform raw ideas into validated mini-startup concepts using local LLMs",
    version="1.0.0",
    lifespan=lifespan
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/")
async def root():
    return {
        "message": "AI Innovation Hub API",
        "docs": "/docs",
        "status": "running"
    }


@app.get("/health")
async def health_check():
    llm_status = await check_llm_status()
    return {
        "api": "healthy",
        "llm": llm_status,
    }


# --- Lightweight chat endpoint (used by the frontend floating Q/A bot) ---


def _clean_chat_reply(text: str) -> str:
    """Best-effort cleanup of markdown-ish formatting for chat UI.

    The goal is to prevent noisy outputs like **bold**, `code`, and other symbols.
    """

    raw = (text or "").strip()
    if not raw:
        return ""

    # Remove fenced code blocks but keep their contents.
    raw = re.sub(r"```(?:[a-zA-Z0-9_-]+)?\n([\s\S]*?)```", r"\1", raw)

    cleaned_lines: list[str] = []
    for line in raw.splitlines():
        s = line.rstrip()
        # strip heading markers
        s = re.sub(r"^\s{0,3}#{1,6}\s+", "", s)
        # strip common markdown bullets (keep indentation)
        s = re.sub(r"^(\s*)[-*+]\s+", r"\1", s)
        # strip emphasis/backticks/links
        s = clean_inline_markdown(s)
        cleaned_lines.append(s)

    out = "\n".join(cleaned_lines)
    # Collapse excessive blank lines
    out = re.sub(r"\n{3,}", "\n\n", out).strip()
    return out


@app.post("/chat", response_model=ChatResponse)
async def chat(req: ChatRequest):
    """Simple chat + generation endpoint.

    - Uses the backend primary model (defaults to mistral via env).
    - Accepts normal text messages (not only Q/A).
    - `context` is optional and can include UI state (current idea, selected project, etc.).
    """

    message = (req.message or "").strip()
    if not message:
        raise HTTPException(status_code=400, detail="Message is required")

    # Fast-path: greetings should feel natural and short (avoid dumping app descriptions).
    lowered = message.lower().strip()
    if re.fullmatch(r"(hi|hello|hey|hii+|hai|yo|good\s+morning|good\s+afternoon|good\s+evening)(!|\.|\s+there|\s+ai|\s+bot|\s+buddy)?", lowered):
        return ChatResponse(reply="Hi! How can I help?")

    ctx = (req.context or "").strip()
    prompt = (
        "You are a helpful assistant inside AI Innovation Hub.\n"
        "Respond naturally to normal messages.\n"
        "Be concise by default (1-2 short sentences). Only give long explanations if the user asks.\n"
        "Do not list app features unless the user requests them.\n"
        "When the user asks to generate a project description, output ONLY the description text (no bullets, no markdown).\n\n"
    )
    if ctx:
        prompt += f"Context:\n{ctx}\n\n"
    prompt += f"User: {message}\nAssistant:"

    try:
        reply = await call_primary_model(prompt, system_prompt=SYSTEM_PROMPT)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    return ChatResponse(reply=_clean_chat_reply(reply))


# --- Project Endpoints ---

@app.post("/projects", response_model=ProjectResponse)
def create_project(
    project_data: ProjectCreate,
    session: Session = Depends(get_session)
):
    """Create a new project and initialize all phases."""
    project = Project(**project_data.model_dump())
    session.add(project)
    session.commit()
    session.refresh(project)
    
    # Create all phases for this project
    create_project_phases(session, project)
    
    return project


@app.get("/projects", response_model=List[ProjectResponse])
def list_projects(session: Session = Depends(get_session)):
    """List all projects."""
    projects = session.exec(select(Project).order_by(Project.created_at.desc())).all()
    return projects


@app.get("/projects/{project_id}", response_model=ProjectDetailResponse)
def get_project(project_id: int, session: Session = Depends(get_session)):
    """Get project with all phases and artifacts."""
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Ensure new phases exist for older projects/databases
    phases = ensure_project_phases(session, project)
    
    phase_responses = []
    for phase in phases:
        artifacts = session.exec(
            select(Artifact)
            .where(Artifact.phase_id == phase.id)
            .order_by(Artifact.artifact_type)
        ).all()
        phase_responses.append(PhaseResponse(
            id=phase.id,
            project_id=phase.project_id,
            phase_type=phase.phase_type,
            status=phase.status,
            started_at=phase.started_at,
            completed_at=phase.completed_at,
            artifacts=[ArtifactResponse.model_validate(a) for a in artifacts]
        ))
    
    return ProjectDetailResponse(
        id=project.id,
        name=project.name,
        description=project.description,
        user_type=project.user_type,
        constraints=project.constraints,
        skills=project.skills,
        time_available=project.time_available,
        created_at=project.created_at,
        updated_at=project.updated_at,
        phases=phase_responses
    )


@app.delete("/projects/{project_id}")
def delete_project(project_id: int, session: Session = Depends(get_session)):
    """Delete a project and all its phases/artifacts."""
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    
    # Delete artifacts first
    phases = session.exec(select(Phase).where(Phase.project_id == project_id)).all()
    for phase in phases:
        artifacts = session.exec(select(Artifact).where(Artifact.phase_id == phase.id)).all()
        for artifact in artifacts:
            session.delete(artifact)
        session.delete(phase)
    
    session.delete(project)
    session.commit()
    
    return {"message": "Project deleted successfully"}


# --- Phase Endpoints ---

@app.get("/projects/{project_id}/phases", response_model=List[PhaseResponse])
def list_phases(project_id: int, session: Session = Depends(get_session)):
    """List all phases for a project."""
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    phases = ensure_project_phases(session, project)
    
    phase_responses = []
    for phase in phases:
        artifacts = session.exec(
            select(Artifact).where(Artifact.phase_id == phase.id)
        ).all()
        phase_responses.append(PhaseResponse(
            id=phase.id,
            project_id=phase.project_id,
            phase_type=phase.phase_type,
            status=phase.status,
            started_at=phase.started_at,
            completed_at=phase.completed_at,
            artifacts=[ArtifactResponse.model_validate(a) for a in artifacts]
        ))
    
    return phase_responses


@app.post("/projects/{project_id}/phases/{phase_type}/generate", response_model=GenerateResponse)
async def generate_phase(
    project_id: int,
    phase_type: PhaseType,
    request: GenerateRequest = GenerateRequest(),
    session: Session = Depends(get_session)
):
    """Generate artifacts for a specific phase using local LLM."""
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    
    # Ensure new phases exist for older projects/databases
    ensure_project_phases(session, project)

    phase = get_phase_by_type(session, project_id, phase_type)
    if not phase:
        raise HTTPException(status_code=404, detail="Phase not found")
    
    # Generate artifacts
    artifacts = await generate_phase_artifacts(
        session=session,
        project=project,
        phase=phase,
        regenerate=request.regenerate,
        custom_instructions=request.custom_instructions
    )
    
    # Refresh phase to get updated status
    session.refresh(phase)
    
    return GenerateResponse(
        phase=PhaseResponse(
            id=phase.id,
            project_id=phase.project_id,
            phase_type=phase.phase_type,
            status=phase.status,
            started_at=phase.started_at,
            completed_at=phase.completed_at,
            artifacts=[ArtifactResponse.model_validate(a) for a in artifacts]
        ),
        artifacts=[ArtifactResponse.model_validate(a) for a in artifacts],
        message=f"Generated {len(artifacts)} artifacts for {phase_type.value} phase"
    )


# --- Artifact Endpoints ---

@app.get("/projects/{project_id}/artifacts", response_model=List[ArtifactResponse])
def list_artifacts(project_id: int, session: Session = Depends(get_session)):
    """List all artifacts for a project."""
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    
    artifacts = get_all_artifacts(session, project_id)
    return artifacts


@app.get("/artifacts/{artifact_id}", response_model=ArtifactResponse)
def get_artifact(artifact_id: int, session: Session = Depends(get_session)):
    """Get a specific artifact."""
    artifact = session.get(Artifact, artifact_id)
    if not artifact:
        raise HTTPException(status_code=404, detail="Artifact not found")
    return artifact


# --- Export Endpoints ---

@app.get("/projects/{project_id}/export")
def export_project(project_id: int, session: Session = Depends(get_session)):
    """Export all project artifacts as a single markdown document.

    This is primarily aimed at Notion-friendly markdown exports.
    """
    project, markdown = build_project_markdown(session, project_id)

    return {
        "project_id": project_id,
        "format": "markdown",
        "content": markdown,
    }


@app.get("/projects/{project_id}/export-file")
def export_project_file(project_id: int, format: str = "pdf", session: Session = Depends(get_session)):
    """Export full project spec as a downloadable file (markdown/PDF/DOCX)."""
    project, markdown = build_project_markdown(session, project_id)
    safe_name = project.name.replace(" ", "_")
    base_filename = f"{safe_name}_spec"

    fmt = format.lower()

    if fmt == "markdown":
        buffer = BytesIO(markdown.encode("utf-8"))
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="text/markdown",
            headers={
                "Content-Disposition": f'attachment; filename="{base_filename}.md"',
            },
        )
    if fmt == "pdf":
        return markdown_to_pdf(markdown, base_filename)
    if fmt == "docx":
        return markdown_to_docx(markdown, base_filename)

    raise HTTPException(status_code=400, detail="Unsupported export format. Use 'markdown', 'pdf', or 'docx'.")


@app.post("/projects/{project_id}/export-custom")
def export_project_custom(
    project_id: int,
    request: ExportCustomRequest,
    session: Session = Depends(get_session),
):
    """Export a "final" spec using selected variants where provided, and
    latest versions for all other artifact types.

    - request.artifact_ids: explicit overrides (usually one per artifact_type)
    - all other artifact types: use the latest version available for that type
    """
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Get all artifacts for this project (ordered by phase + type)
    all_artifacts = get_all_artifacts(session, project_id)
    if not all_artifacts:
        raise HTTPException(status_code=400, detail="No artifacts found for this project")

    override_ids = set(request.artifact_ids or [])

    # Step 1: decide which artifact to use per artifact_type
    chosen_by_type: dict[str, Artifact] = {}
    override_types: set[str] = set()

    for artifact in all_artifacts:
        key = artifact.artifact_type.value

        # If this artifact is explicitly selected, prefer it for this type
        if artifact.id in override_ids:
            chosen_by_type[key] = artifact
            override_types.add(key)
            continue

        # If this type already has an override, skip other versions
        if key in override_types:
            continue

        # For non-overridden types, keep the latest version
        existing = chosen_by_type.get(key)
        if existing is None or artifact.version > existing.version:
            chosen_by_type[key] = artifact

    # Step 2: build a list of chosen artifacts preserving phase/type ordering
    chosen_ids = {a.id for a in chosen_by_type.values()}
    ordered_chosen = [a for a in all_artifacts if a.id in chosen_ids]

    # Build markdown similar to the full export, but only for the chosen artifacts
    markdown = f"# {project.name} - Final Spec\n\n"
    markdown += f"**Description:** {project.description}\n\n"
    markdown += f"**User Type:** {project.user_type.value}\n"
    markdown += f"**Constraints:** {project.constraints or 'None'}\n"
    markdown += f"**Skills:** {project.skills or 'Not specified'}\n"
    markdown += f"**Time Available:** {project.time_available or 'Not specified'}\n\n"
    markdown += "---\n\n"

    current_phase = None
    for artifact in ordered_chosen:
        phase = session.get(Phase, artifact.phase_id)
        if phase and phase.phase_type != current_phase:
            current_phase = phase.phase_type
            markdown += f"## {current_phase.value.replace('_', ' ').title()}\n\n"

        markdown += f"### {artifact.artifact_type.value.replace('_', ' ').title()} (v{artifact.version})\n\n"
        markdown += artifact.content + "\n\n"
        markdown += "---\n\n"

    safe_name = project.name.replace(" ", "_")
    base_filename = f"{safe_name}_final_spec"

    fmt = request.format.lower()

    if fmt == "markdown":
        return {
            "project_id": project_id,
            "format": "markdown",
            "content": markdown,
        }
    if fmt == "pdf":
        return markdown_to_pdf(markdown, base_filename)
    if fmt == "docx":
        return markdown_to_docx(markdown, base_filename)

    raise HTTPException(status_code=400, detail="Unsupported export format. Use 'markdown', 'pdf', or 'docx'.")


# --- File-by-file CodeGen Endpoints ---


@app.get("/projects/{project_id}/codegen/sessions/latest", response_model=Optional[CodeGenSessionResponse])
def get_latest_codegen_session(project_id: int, session: Session = Depends(get_session)):
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    statement = (
        select(CodeGenSession)
        .where(CodeGenSession.project_id == project_id)
        .order_by(CodeGenSession.created_at.desc())
    )
    s = session.exec(statement).first()
    return s


@app.post("/projects/{project_id}/codegen/sessions", response_model=CodeGenSessionResponse)
def create_codegen_session(
    project_id: int,
    req: CodeGenCreateSessionRequest = CodeGenCreateSessionRequest(),
    session: Session = Depends(get_session),
):
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    model_used = req.model_used or "coder"
    cg = CodeGenSession(project_id=project_id, model_used=model_used)
    session.add(cg)
    session.commit()
    session.refresh(cg)
    return cg


@app.get("/codegen/sessions/{session_id}", response_model=CodeGenSessionDetailResponse)
def get_codegen_session(session_id: int, session: Session = Depends(get_session)):
    cg = session.get(CodeGenSession, session_id)
    if not cg:
        raise HTTPException(status_code=404, detail="CodeGen session not found")

    files = session.exec(
        select(CodeGenFile).where(CodeGenFile.session_id == session_id).order_by(CodeGenFile.path)
    ).all()
    messages = session.exec(
        select(CodeGenMessage).where(CodeGenMessage.session_id == session_id).order_by(CodeGenMessage.created_at)
    ).all()

    return CodeGenSessionDetailResponse(
        id=cg.id,
        project_id=cg.project_id,
        status=cg.status,
        model_used=cg.model_used,
        created_at=cg.created_at,
        updated_at=cg.updated_at,
        files=files,
        messages=messages,
    )


@app.put("/codegen/sessions/{session_id}/files", response_model=CodeGenFileResponse)
def upsert_codegen_file(
    session_id: int,
    req: CodeGenUpsertFileRequest,
    session: Session = Depends(get_session),
):
    cg = session.get(CodeGenSession, session_id)
    if not cg:
        raise HTTPException(status_code=404, detail="CodeGen session not found")

    try:
        safe_path = safe_relpath(req.path)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

    existing = session.exec(
        select(CodeGenFile)
        .where(CodeGenFile.session_id == session_id)
        .where(CodeGenFile.path == safe_path)
    ).first()

    from datetime import datetime

    if existing:
        existing.content = req.content
        existing.updated_at = datetime.utcnow()
        session.add(existing)
        session.commit()
        session.refresh(existing)
        return existing

    f = CodeGenFile(session_id=session_id, path=safe_path, content=req.content)
    session.add(f)
    session.commit()
    session.refresh(f)
    return f


@app.get("/codegen/sessions/{session_id}/plan", response_model=CodeGenPlanResponse)
def get_codegen_plan(session_id: int, session: Session = Depends(get_session)):
    cg = session.get(CodeGenSession, session_id)
    if not cg:
        raise HTTPException(status_code=404, detail="CodeGen session not found")

    project = session.get(Project, cg.project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Try to infer tech stack from stored artifacts
    artifacts = session.exec(
        select(Artifact)
        .join(Phase)
        .where(Phase.project_id == project.id)
        .order_by(Artifact.created_at)
    ).all()

    latest: dict[str, Artifact] = {}
    for a in artifacts:
        k = a.artifact_type.value
        cur = latest.get(k)
        if cur is None or a.version >= cur.version:
            latest[k] = a

    tech_stack = latest.get(ArtifactType.TECH_STACK.value).content if latest.get(ArtifactType.TECH_STACK.value) else None
    suggested = suggest_file_paths(tech_stack)

    return CodeGenPlanResponse(session_id=session_id, suggested_paths=suggested)


@app.post("/codegen/sessions/{session_id}/plan/apply", response_model=CodeGenApplyPlanResponse)
def apply_codegen_plan(session_id: int, session: Session = Depends(get_session)):
    cg = session.get(CodeGenSession, session_id)
    if not cg:
        raise HTTPException(status_code=404, detail="CodeGen session not found")

    project = session.get(Project, cg.project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    plan = get_codegen_plan(session_id, session)

    existing = session.exec(select(CodeGenFile).where(CodeGenFile.session_id == session_id)).all()
    existing_paths = {f.path for f in existing}

    from datetime import datetime

    created = 0
    for p in plan.suggested_paths:
        sp = safe_relpath(p)
        if sp in existing_paths:
            continue
        f = CodeGenFile(session_id=session_id, path=sp, content="")
        f.updated_at = datetime.utcnow()
        session.add(f)
        created += 1

    if created:
        cg.updated_at = datetime.utcnow()
        session.add(cg)
        session.commit()

    total = session.exec(select(CodeGenFile).where(CodeGenFile.session_id == session_id)).all()
    return CodeGenApplyPlanResponse(session_id=session_id, created_files=created, total_files=len(total))


@app.delete("/codegen/sessions/{session_id}/files/{file_id}")
def delete_codegen_file(session_id: int, file_id: int, session: Session = Depends(get_session)):
    cg = session.get(CodeGenSession, session_id)
    if not cg:
        raise HTTPException(status_code=404, detail="CodeGen session not found")

    f = session.get(CodeGenFile, file_id)
    if not f or f.session_id != session_id:
        raise HTTPException(status_code=404, detail="File not found")

    session.delete(f)
    session.commit()
    return {"message": "File deleted"}


@app.post("/codegen/sessions/{session_id}/generate-file", response_model=CodeGenFileResponse)
async def generate_codegen_file(
    session_id: int,
    req: CodeGenGenerateFileRequest,
    session: Session = Depends(get_session),
):
    settings = get_settings()
    cg = session.get(CodeGenSession, session_id)
    if not cg:
        raise HTTPException(status_code=404, detail="CodeGen session not found")

    project = session.get(Project, cg.project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Gather latest artifacts that help generation (best-effort)
    artifacts = session.exec(
        select(Artifact)
        .join(Phase)
        .where(Phase.project_id == project.id)
        .order_by(Artifact.created_at)
    ).all()

    latest: dict[str, Artifact] = {}
    for a in artifacts:
        k = a.artifact_type.value
        cur = latest.get(k)
        if cur is None or a.version >= cur.version:
            latest[k] = a

    tech_stack = latest.get(ArtifactType.TECH_STACK.value).content if latest.get(ArtifactType.TECH_STACK.value) else None
    api_design = latest.get(ArtifactType.API_DESIGN.value).content if latest.get(ArtifactType.API_DESIGN.value) else None
    feature_list = latest.get(ArtifactType.FEATURE_LIST.value).content if latest.get(ArtifactType.FEATURE_LIST.value) else None
    user_flows = latest.get(ArtifactType.USER_FLOWS.value).content if latest.get(ArtifactType.USER_FLOWS.value) else None

    try:
        target_path = safe_relpath(req.path)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

    existing_files = session.exec(
        select(CodeGenFile).where(CodeGenFile.session_id == session_id).order_by(CodeGenFile.path)
    ).all()
    existing_paths = [f.path for f in existing_files]

    prompt = build_generate_file_prompt(
        project_name=project.name,
        project_description=project.description,
        tech_stack=tech_stack,
        api_design=api_design,
        feature_list=feature_list,
        user_flows=user_flows,
        existing_files=existing_paths,
        target_path=target_path,
        instructions=req.instructions,
    )

    # Store user message
    session.add(CodeGenMessage(session_id=session_id, role="user", content=f"Generate {target_path}\n\n{req.instructions or ''}".strip()))
    session.commit()

    try:
        output = await call_coder_model(prompt, SYSTEM_PROMPT)
        if settings.enable_guardrails:
            # Validate/normalize JSON output in a consistent way
            sf = validate_single_file_payload(output, max_file_chars=settings.max_file_chars)
            parsed = {"path": sf.normalized_path(), "content": sf.content}
        else:
            parsed = parse_file_json(output)
    except (GuardrailError, Exception) as e:
        # Store assistant output for debugging
        session.add(CodeGenMessage(session_id=session_id, role="assistant", content=f"[Generation failed]\n{str(e)}"))
        session.commit()
        raise HTTPException(status_code=400, detail=f"Could not generate/parse file JSON: {str(e)}")

    # Upsert file
    from datetime import datetime

    existing = session.exec(
        select(CodeGenFile)
        .where(CodeGenFile.session_id == session_id)
        .where(CodeGenFile.path == parsed["path"])
    ).first()

    if existing:
        existing.content = parsed["content"]
        existing.updated_at = datetime.utcnow()
        session.add(existing)
        session.commit()
        session.refresh(existing)
        f = existing
    else:
        f = CodeGenFile(session_id=session_id, path=parsed["path"], content=parsed["content"])
        session.add(f)
        session.commit()
        session.refresh(f)

    # Store assistant message (raw JSON output)
    session.add(CodeGenMessage(session_id=session_id, role="assistant", content=output))
    cg.updated_at = datetime.utcnow()
    session.add(cg)
    session.commit()

    return f


@app.get("/codegen/sessions/{session_id}/zip")
def download_codegen_zip(session_id: int, session: Session = Depends(get_session)):
    cg = session.get(CodeGenSession, session_id)
    if not cg:
        raise HTTPException(status_code=404, detail="CodeGen session not found")

    files = session.exec(select(CodeGenFile).where(CodeGenFile.session_id == session_id)).all()
    if not files:
        raise HTTPException(status_code=400, detail="No files in this codegen session")

    project = session.get(Project, cg.project_id)
    base_filename = f"{(project.name if project else 'project').replace(' ', '_')}_code"

    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for f in files:
            rel = safe_relpath(f.path)
            zf.writestr(rel, normalize_content(f.content))

    zip_buffer.seek(0)
    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={
            "Content-Disposition": f'attachment; filename="{base_filename}.zip"',
        },
    )


@app.post("/projects/{project_id}/codezip")
def export_project_code_zip(
    project_id: int,
    request: CodeZipRequest = CodeZipRequest(),
    session: Session = Depends(get_session),
):
    settings = get_settings()
    """Generate a ZIP file from the latest code_generation artifact (strict JSON output).

    The code_generation artifact content must be JSON:
    { files: [{ path, content }], instructions: "..." }

    If request.artifact_id is provided, that artifact will be used instead.
    """
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Ensure code_generation phase exists
    ensure_project_phases(session, project)

    # Find artifact to use
    artifact: Artifact | None = None
    if request.artifact_id is not None:
        artifact = session.get(Artifact, request.artifact_id)
        if not artifact:
            raise HTTPException(status_code=404, detail="Artifact not found")
    else:
        # Latest generated_app artifact for this project
        statement = (
            select(Artifact)
            .join(Phase)
            .where(Phase.project_id == project_id)
            .where(Phase.phase_type == PhaseType.CODE_GENERATION)
            .where(Artifact.artifact_type == ArtifactType.GENERATED_APP)
            .order_by(Artifact.version.desc())
        )
        artifact = session.exec(statement).first()

    if not artifact:
        raise HTTPException(status_code=400, detail="No generated code found. Run the code_generation phase first.")

    # Parse JSON output - handle LLM responses that may be wrapped in markdown code blocks
    def extract_json_from_content(content: str) -> dict:
        """Extract JSON from content that may be wrapped in markdown code blocks."""
        import re
        
        def try_parse_json(text: str) -> dict | None:
            """Try to parse JSON, with fallback to fix common LLM formatting issues."""
            raw = (text or "").strip()

            # Common failure: older prompts asked the model to wrap JSON in double braces: {{...}}
            # Strip one brace from each side if it looks like that.
            if raw.startswith("{{") and raw.endswith("}}"):
                raw = raw[1:-1].strip()

            # Try direct parse first
            try:
                return json.loads(raw)
            except json.JSONDecodeError:
                pass
            
            # Try to fix unescaped newlines in string values
            # This is a more sophisticated approach that handles multiline strings
            try:
                # Replace actual newlines with escaped newlines, but be smart about it
                # First, remove any \r characters
                fixed = text.replace('\r\n', '\n').replace('\r', '\n')
                
                # Try to fix newlines that appear within JSON string values
                # This regex finds content between quotes and escapes newlines within
                def escape_newlines_in_strings(match):
                    s = match.group(0)
                    # Replace actual newlines with \n escape sequence
                    inner = s[1:-1]  # Remove surrounding quotes
                    inner = inner.replace('\n', '\\n')
                    inner = inner.replace('\t', '\\t')
                    return '"' + inner + '"'
                
                # Match JSON strings (handling escaped quotes)
                fixed = re.sub(r'"(?:[^"\\]|\\.)*"', escape_newlines_in_strings, fixed, flags=re.DOTALL)
                return json.loads(fixed)
            except (json.JSONDecodeError, Exception):
                pass
            
            return None
        
        def extract_files_manually(text: str) -> list:
            """Manually extract file objects from malformed JSON."""
            files = []
            
            # Pattern to find file path
            path_pattern = r'"path"\s*:\s*"([^"]+)"'
            
            # Find all paths first
            path_matches = list(re.finditer(path_pattern, text))
            
            for i, path_match in enumerate(path_matches):
                path = path_match.group(1)
                start_pos = path_match.end()
                
                # Find the content field after this path
                content_search = text[start_pos:start_pos + 50]  # Look ahead a bit
                content_marker = re.search(r'"content"\s*:\s*"', content_search)
                
                if content_marker:
                    content_start = start_pos + content_marker.end()
                    
                    # Find where this file's content ends
                    # Look for the closing quote that's not escaped
                    pos = content_start
                    escaped = False
                    content_chars = []
                    
                    while pos < len(text):
                        ch = text[pos]
                        if escaped:
                            content_chars.append(ch)
                            escaped = False
                        elif ch == '\\':
                            content_chars.append(ch)
                            escaped = True
                        elif ch == '"':
                            # End of content string
                            break
                        else:
                            content_chars.append(ch)
                        pos += 1
                    
                    file_content = ''.join(content_chars)
                    
                    # Unescape the content
                    file_content = file_content.replace('\\n', '\n')
                    file_content = file_content.replace('\\t', '\t')
                    file_content = file_content.replace('\\"', '"')
                    file_content = file_content.replace('\\\\', '\\')
                    
                    if path and file_content:
                        files.append({"path": path, "content": file_content})
            
            return files
        
        # Step 1: Try direct JSON parse
        result = try_parse_json(content)
        if result and result.get("files"):
            return result
        
        # Step 2: Remove markdown code blocks and try again
        clean_content = content
        code_block_pattern = r'```(?:json)?\s*([\s\S]*?)```'
        matches = re.findall(code_block_pattern, content)
        for match in matches:
            result = try_parse_json(match.strip())
            if result and result.get("files"):
                return result
            clean_content = match.strip()
        
        # Step 3: Find JSON object boundaries and try parsing
        # Look for outermost { ... } that contains "files"
        if '"files"' in clean_content:
            start = clean_content.find('{')
            if start != -1:
                # Find matching closing brace
                depth = 0
                for i, ch in enumerate(clean_content[start:], start):
                    if ch == '{':
                        depth += 1
                    elif ch == '}':
                        depth -= 1
                        if depth == 0:
                            json_str = clean_content[start:i+1]
                            result = try_parse_json(json_str)
                            if result and result.get("files"):
                                return result
                            break
        
        # Step 4: Manual extraction as last resort
        files = extract_files_manually(content)
        if files:
            # Try to find instructions
            instr_match = re.search(r'"instructions"\s*:\s*"((?:[^"\\]|\\.)*)"', content)
            instructions = instr_match.group(1) if instr_match else "See README.md for setup instructions"
            instructions = instructions.replace('\\n', '\n').replace('\\"', '"')
            return {"files": files, "instructions": instructions}
        
        raise ValueError("Could not extract valid file data from LLM response")
    
    try:
        if settings.enable_guardrails:
            payload = validate_generated_app_payload(
                artifact.content,
                max_files=settings.max_generated_files,
                max_file_chars=settings.max_file_chars,
                max_total_chars=settings.max_total_chars,
            )
            # Convert to a plain dict the rest of the handler expects
            files = [{"path": f.normalized_path(), "content": f.content} for f in payload.files]
        else:
            payload = extract_json_from_content(artifact.content)
            files = payload.get("files")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not parse/validate generated JSON. Error: {str(e)}")

    if not isinstance(files, list) or len(files) == 0:
        # Check if the response indicates missing context
        if "Not yet generated" in str(getattr(payload, "model_dump", lambda: payload)()):
            raise HTTPException(
                status_code=400,
                detail="Code generation failed: Previous phases (ideation, concept_design, tech_build) must be completed first. Please generate those phases before running code generation.",
            )
        raise HTTPException(status_code=400, detail="Generated code JSON must contain a non-empty 'files' array. Try regenerating the code generation phase.")

    def _safe_relpath(p: str) -> str:
        p = p.replace("\\", "/")
        p = p.lstrip("/")
        # Prevent path traversal
        norm = os.path.normpath(p).replace("\\", "/")
        if norm.startswith("../") or norm == ".." or ":" in norm:
            raise HTTPException(status_code=400, detail=f"Unsafe file path in generated output: {p}")
        return norm

    base_filename = f"{project.name.replace(' ', '_')}_code"

    # Write to temp folder then zip
    with tempfile.TemporaryDirectory(prefix="ai_innovation_hub_") as tmpdir:
        root_dir = os.path.join(tmpdir, base_filename)
        os.makedirs(root_dir, exist_ok=True)

        def normalize_content(content: str) -> str:
            """Normalize file content - convert escaped sequences to actual characters."""
            # Convert escaped newlines to real newlines
            content = content.replace('\\n', '\n')
            # Convert escaped tabs to real tabs
            content = content.replace('\\t', '\t')
            # Convert escaped quotes
            content = content.replace('\\"', '"')
            # Convert escaped backslashes (do this last)
            content = content.replace('\\\\', '\\')
            # Remove any carriage returns for consistent line endings
            content = content.replace('\r\n', '\n').replace('\r', '\n')
            return content

        for f in files:
            if not isinstance(f, dict):
                continue
            rel_path = f.get("path")
            content = f.get("content")
            if not isinstance(rel_path, str) or not isinstance(content, str):
                continue

            # Normalize the content to have proper formatting
            content = normalize_content(content)

            safe_path = _safe_relpath(rel_path)
            abs_path = os.path.join(root_dir, safe_path)
            os.makedirs(os.path.dirname(abs_path), exist_ok=True)
            with open(abs_path, "w", encoding="utf-8", newline="\n") as fp:
                fp.write(content)

        zip_buffer = BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for folder, _, filenames in os.walk(root_dir):
                for filename in filenames:
                    abs_path = os.path.join(folder, filename)
                    arcname = os.path.relpath(abs_path, root_dir)
                    zf.write(abs_path, arcname=arcname)

        zip_buffer.seek(0)

        return StreamingResponse(
            zip_buffer,
            media_type="application/zip",
            headers={
                "Content-Disposition": f'attachment; filename="{base_filename}.zip"',
            },
        )


@app.get("/artifacts/{artifact_id}/export")
def export_artifact(artifact_id: int, format: str = "markdown", session: Session = Depends(get_session)):
    """Export a single artifact in markdown, PDF, or DOCX format.

    - markdown: JSON response with a markdown string (Notion-friendly)
    - pdf:      application/pdf binary stream
    - docx:     application/vnd.openxmlformats-officedocument.wordprocessingml.document binary stream
    """
    artifact = session.get(Artifact, artifact_id)
    if not artifact:
        raise HTTPException(status_code=404, detail="Artifact not found")

    phase = session.get(Phase, artifact.phase_id)
    project = session.get(Project, phase.project_id) if phase else None

    project_name = (project.name if project else "project").replace(" ", "_")
    artifact_name = artifact.artifact_type.value.replace(" ", "_")
    base_filename = f"{project_name}_{artifact_name}_v{artifact.version}"

    header = f"{project.name if project else 'Project'} - {artifact.artifact_type.value.replace('_', ' ').title()} (v{artifact.version})"
    markdown = f"# {header}\n\n{artifact.content}\n"

    fmt = format.lower()

    if fmt == "markdown":
        return {
            "artifact_id": artifact_id,
            "format": "markdown",
            "content": markdown,
        }
    if fmt == "pdf":
        return markdown_to_pdf(markdown, base_filename)
    if fmt == "docx":
        return markdown_to_docx(markdown, base_filename)

    raise HTTPException(status_code=400, detail="Unsupported export format. Use 'markdown', 'pdf', or 'docx'.")
